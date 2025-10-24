#!/usr/bin/env python3
"""
Import Glossary Script for TRACE
Reads glossary files (Word documents with tables) and imports definitions into knowledge base.

Handles multiple glossary types:
  - System Groups: =10, =61, etc. (system categories, not locations)
  - Terminals: -X10, -X11, etc. (terminal types and wire types)
  - Locations: (future) actual physical locations

Expected format: 2-column tables
  Table 1 - System Groups:
    Column 1: Code (e.g., =10, =61)
    Column 2: Definition (e.g., MV-Supply, PLC/CMS)

  Table 2 - Terminals:
    Column 1: Terminal Code (e.g., -X10, -X11)
    Column 2: Description (e.g., Power terminal)

Usage:
  python3 import_glossary.py glossary/electrical_sections.docx
"""

import json
import sys
import os
from datetime import datetime

# Try importing Word document reader
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("⚠ python-docx not available. Install with: pip3 install python-docx")


class GlossaryImporter:
    """Imports glossary definitions into TRACE knowledge base"""

    def __init__(self, knowledge_file="trace_knowledge.json"):
        self.knowledge_file = knowledge_file
        self.knowledge = self._load_knowledge()

    def _load_knowledge(self):
        """Load existing knowledge base"""
        if os.path.exists(self.knowledge_file):
            with open(self.knowledge_file, 'r') as f:
                return json.load(f)
        else:
            print(f"✗ Knowledge file not found: {self.knowledge_file}")
            sys.exit(1)

    def _save_knowledge(self):
        """Save updated knowledge base"""
        with open(self.knowledge_file, 'w') as f:
            json.dump(self.knowledge, f, indent=2)
        print(f"✓ Knowledge saved to {self.knowledge_file}")

    def _ensure_glossary_sections(self):
        """Ensure glossary sections exist in knowledge base"""
        created = []
        if "system_groups" not in self.knowledge:
            self.knowledge["system_groups"] = {}
            created.append("system_groups")
        if "terminals" not in self.knowledge:
            self.knowledge["terminals"] = {}
            created.append("terminals")
        if created:
            print(f"✓ Created sections: {', '.join(created)}")

    def _add_definition(self, code, definition, source):
        """Add a definition to the appropriate section based on code format"""
        if code.startswith('='):
            # System Group (e.g., =10, =61)
            self.knowledge["system_groups"][code] = {
                "code": code,
                "definition": definition,
                "source": source,
                "added": datetime.now().isoformat()
            }
            return "system_group"
        elif code.startswith('-'):
            # Terminal (e.g., -X10, -X11)
            self.knowledge["terminals"][code] = {
                "code": code,
                "terminal_type": definition,
                "source": source,
                "added": datetime.now().isoformat()
            }
            return "terminal"
        else:
            # Unknown format
            print(f"  ⚠ Unknown format: {code} (expected = or - prefix)")
            return None

    def import_from_docx(self, doc_path):
        """Import glossary from Word document (.docx format)"""
        if not DOCX_AVAILABLE:
            print("✗ python-docx library required for .docx files")
            print("  Install with: pip3 install python-docx")
            return False

        try:
            doc = Document(doc_path)
        except Exception as e:
            print(f"✗ Could not open {doc_path}")
            print(f"  Error: {e}")
            print("\n💡 Note: This script works with .docx files (newer format).")
            print("   If you have a .doc file (old format), please:")
            print("   1. Open it in Word")
            print("   2. Save As -> .docx format")
            print("   3. Run this script again")
            return False

        print(f"\n📄 Reading {doc_path}...")

        self._ensure_glossary_sections()

        system_groups_added = 0
        terminals_added = 0

        # Check if document has tables
        if doc.tables:
            print(f"   Found {len(doc.tables)} table(s)\n")
            # Process ALL tables in the document
            for table_num, table in enumerate(doc.tables, start=1):
                print(f"\n--- Processing Table {table_num} ---")

                # Skip header row, process data rows
                for i, row in enumerate(table.rows[1:], start=1):
                    cells = row.cells
                    if len(cells) >= 2:
                        code = cells[0].text.strip()
                        definition = cells[1].text.strip()

                        if code and definition:
                            added = self._add_definition(code, definition, doc_path)
                            if added == "system_group":
                                system_groups_added += 1
                                print(f"  ✓ System Group: {code:10} → {definition}")
                            elif added == "terminal":
                                terminals_added += 1
                                print(f"  ✓ Terminal:     {code:10} → {definition}")
        else:
            # No tables found - try reading as plain text with spacing
            print("   No tables found - reading as plain text\n")

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Skip header lines
                if text.lower().startswith('system') or text.lower().startswith('terminal'):
                    print(f"\n--- {text} ---")
                    continue

                # Try to parse line with multiple spaces/tabs between code and definition
                # Split on 2+ spaces or tabs
                import re
                parts = re.split(r'\s{2,}|\t+', text, maxsplit=1)

                if len(parts) >= 2:
                    code = parts[0].strip()
                    definition = parts[1].strip()

                    if code and definition:
                        added = self._add_definition(code, definition, doc_path)
                        if added == "system_group":
                            system_groups_added += 1
                            print(f"  ✓ System Group: {code:10} → {definition}")
                        elif added == "terminal":
                            terminals_added += 1
                            print(f"  ✓ Terminal:     {code:10} → {definition}")
                elif text.startswith('=') or text.startswith('-'):
                    # Single word entries or malformed - show warning
                    print(f"  ⚠ Could not parse: {text[:50]}")

        self._save_knowledge()
        print(f"\n{'='*60}")
        print(f"✓ Import Complete!")
        print(f"  System Groups: {system_groups_added}")
        print(f"  Terminals:     {terminals_added}")
        print(f"  Total:         {system_groups_added + terminals_added}")
        print(f"{'='*60}\n")
        return True

    def import_from_manual_input(self):
        """Manually enter definitions one by one"""
        print("\n=== Manual Glossary Entry ===")
        print("Enter definitions in format: CODE | DEFINITION")
        print("Examples:")
        print("  =10 | MV-Supply        (system group)")
        print("  -X10 | Power terminal  (terminal)")
        print("Type 'done' when finished\n")

        self._ensure_glossary_sections()
        system_groups_added = 0
        terminals_added = 0

        while True:
            entry = input("Enter definition (or 'done'): ").strip()

            if entry.lower() == 'done':
                break

            if '|' not in entry:
                print("  ✗ Format should be: CODE | DEFINITION")
                continue

            parts = entry.split('|', 1)
            code = parts[0].strip()
            definition = parts[1].strip()

            if code and definition:
                if code.startswith('='):
                    # System Group
                    self.knowledge["system_groups"][code] = {
                        "code": code,
                        "definition": definition,
                        "source": "manual_entry",
                        "added": datetime.now().isoformat()
                    }
                    system_groups_added += 1
                    print(f"  ✓ Added System Group: {code} → {definition}")
                elif code.startswith('-'):
                    # Terminal
                    self.knowledge["terminals"][code] = {
                        "code": code,
                        "terminal_type": definition,
                        "source": "manual_entry",
                        "added": datetime.now().isoformat()
                    }
                    terminals_added += 1
                    print(f"  ✓ Added Terminal: {code} → {definition}")
                else:
                    print(f"  ✗ Unknown format: {code} (expected = or - prefix)")
            else:
                print("  ✗ Both code and definition required")

        total_added = system_groups_added + terminals_added
        if total_added > 0:
            self._save_knowledge()
            print(f"\n✓ Imported {total_added} definitions successfully!")
            print(f"  System Groups: {system_groups_added}")
            print(f"  Terminals:     {terminals_added}")
        else:
            print("\n  No definitions added.")

        return total_added > 0

    def show_all_definitions(self):
        """Display all glossary definitions"""
        has_data = False

        # Show System Groups
        if "system_groups" in self.knowledge and self.knowledge["system_groups"]:
            print("\n=== System Groups (System Categories) ===\n")
            for code in sorted(self.knowledge["system_groups"].keys()):
                entry = self.knowledge["system_groups"][code]
                print(f"  {code:10} → {entry['definition']}")
            print(f"\n  Total: {len(self.knowledge['system_groups'])} system groups")
            has_data = True

        # Show Terminals
        if "terminals" in self.knowledge and self.knowledge["terminals"]:
            print("\n=== Terminals (Terminal Types) ===\n")
            for code in sorted(self.knowledge["terminals"].keys()):
                entry = self.knowledge["terminals"][code]
                print(f"  {code:10} → {entry['terminal_type']}")
            print(f"\n  Total: {len(self.knowledge['terminals'])} terminals")
            has_data = True

        if not has_data:
            print("\nNo glossary definitions in knowledge base.")
            print("Use: python3 import_glossary.py <file.docx> to import\n")


def main():
    """Main entry point"""
    importer = GlossaryImporter()

    if len(sys.argv) < 2:
        print("Usage: python3 import_glossary.py <path-to-glossary.docx>")
        print("   OR: python3 import_glossary.py --manual")
        print("   OR: python3 import_glossary.py --show")
        sys.exit(1)

    arg = sys.argv[1]

    if arg == "--show":
        importer.show_all_definitions()
    elif arg == "--manual":
        importer.import_from_manual_input()
    else:
        doc_path = arg
        if not os.path.exists(doc_path):
            print(f"✗ File not found: {doc_path}")
            sys.exit(1)

        importer.import_from_docx(doc_path)


if __name__ == "__main__":
    main()
